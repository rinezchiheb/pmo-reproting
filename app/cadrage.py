"""Cadrage du sprint : analyse de début de sprint (estimation, affectation, répartition du RAF) et son rapport PowerPoint.

Trois analyses, mêmes chiffres dans la page /cadrage et dans le PowerPoint :
  1. Macro-chiffrage initial (saisi par l'utilisateur, par FEAT) vs chiffrage actualisé
     = consommé + reste à faire des tâches DEV et QA de la FEAT (avant tout travail : l'estimation des tâches).
     Dépassement : actualisé > macro ; fort dépassement au-delà de +20 %.
  2. Charge de test vs charge de développement, par FEAT : QA ÷ DEV face à l'hypothèse (30 % par défaut).
     Conforme à ±5 points ; « test sous-dimensionné » / « sur-dimensionné » au-delà ; fort écart au-delà de 15 points.
  3. Charge planifiée vs capacité disponible, par ressource : RAF de son profil (DEV ou QA) face à sa capacité prévue
     restante jusqu'au jalon (Tempo : planning − absences − congés − scrum events × affectation).
     Surcharge > 100 % (arbitrage), tendu ≥ 90 %, sous-chargé < 60 % ou sans charge (rééquilibrage possible).
Les bogues (ni DEV ni QA) n'entrent dans aucune charge, comme dans le reste de l'outil.
"""
from __future__ import annotations

import io
from collections import Counter, defaultdict
from typing import Any

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches

from . import dnf, history
from .dnf import (AMBER_PALE, AMBER_TXT, BLUE_PALE, GREEN_DARK, GREEN_PALE, GREY, INK, MUTED, OK_TXT, PANEL, QA_BLUE, RED,
                  RED_PALE, WHITE, _chart, _page_no, _slide, _table, _text, _title, jh, long_date, pct, short_date)

DEFAULT_TEST_PCT = 30.0
MACRO_OVER_RED = 20.0      # % au-delà du macro-chiffrage : fort dépassement
TEST_TOL, TEST_RED = 5.0, 15.0  # points d'écart à l'hypothèse de test : conforme, fort écart
LOAD_HIGH, LOAD_TENSE, LOAD_LOW = 1.0, 0.9, 0.6
SIDES = ("dev", "qa")
PROF = {"dev": "DEV", "qa": "QA"}


def _r(x: float | None, n: int = 2) -> float | None:
    return None if x is None else round(x, n)


# ---------------------------------------------------------------- analyse
def analyse(meta: dict[str, Any], rows: list[dict], tp: dict[str, Any] | None, macro: dict[str, float],
            test_pct: float = DEFAULT_TEST_PCT, tempo_error: str | None = None, chantier: str | None = None) -> dict[str, Any]:
    """macro : FEAT (clé Jira, sinon nom) → macro-chiffrage en JH. tp : données Tempo (/api/tempo), None si indisponible.
    chantier : analyse limitée à ce chantier ; la capacité d'une personne y est sa capacité × sa part d'affectation au chantier."""
    chantiers = sorted({r["chantier"] for r in rows if history.side_of(r) in SIDES}, key=lambda c: (c.startswith("("), c.lower()))
    if chantier:
        rows = [r for r in rows if r["chantier"] == chantier]
    # ---- par FEAT : consommé + reste (à terminaison) et estimé initial, DEV et QA
    feats: dict[str, dict[str, Any]] = {}
    for r in rows:
        sd = history.side_of(r)
        if sd not in SIDES:
            continue
        key = r.get("feat_key") or r.get("feat") or "(Sans FEAT)"
        f = feats.setdefault(key, {"key": key, "name": r.get("feat") or key, "chantiers": Counter(), "n": 0,
                                   **{sd_: {"est": 0.0, "spent": 0.0, "raf": 0.0, "fin": 0.0, "n": 0, "noest": 0} for sd_ in SIDES}})
        f["chantiers"][r["chantier"]] += 1
        f["n"] += 1
        e = f[sd]
        raf = 0.0 if r.get("done") else (r.get("raf_hj") or 0)
        e["n"] += 1
        e["est"] += r.get("estim_hj") or 0
        e["spent"] += r.get("spent_hj") or 0
        e["raf"] += raf
        e["fin"] += (r.get("spent_hj") or 0) + raf
        e["noest"] += not (r.get("estim_hj") or 0) and not r.get("done")
    out_feats = []
    for f in feats.values():
        dev, qa = f["dev"]["fin"], f["qa"]["fin"]
        upd = dev + qa
        m = macro.get(f["key"])
        gap = None if m is None else upd - m
        gpct = None if m is None or m <= 0 else 100 * gap / m
        if m is None:
            ms = {"lvl": "unknown", "word": "macro à saisir"}
        elif gap > 0.005:
            ms = {"lvl": "red" if gpct is None or gpct > MACRO_OVER_RED else "amber", "word": "dépassement"}
        else:
            ms = {"lvl": "green", "word": "dans l'enveloppe"}
        ratio = 100 * qa / dev if dev > 0.005 else None
        theo = test_pct / 100 * dev
        if ratio is None:
            ts = {"lvl": "unknown", "word": "pas de développement"}
        elif abs(ratio - test_pct) <= TEST_TOL:
            ts = {"lvl": "green", "word": "conforme"}
        else:
            ts = {"lvl": "red" if abs(ratio - test_pct) > TEST_RED else "amber",
                  "word": "test sous-dimensionné" if ratio < test_pct else "test sur-dimensionné"}
        out_feats.append({
            "key": f["key"], "name": f["name"], "chantier": f["chantiers"].most_common(1)[0][0], "n": f["n"],
            "dev": {k: _r(v) for k, v in f["dev"].items()}, "qa": {k: _r(v) for k, v in f["qa"].items()},
            "est": _r(f["dev"]["est"] + f["qa"]["est"]), "updated": _r(upd),
            "macro": m, "gap": _r(gap), "gap_pct": _r(gpct, 1), "macro_state": ms,
            "test_ratio": _r(ratio, 1), "test_theo": _r(theo), "test_gap": _r(qa - theo), "test_state": ts,
        })
    out_feats.sort(key=lambda x: (x["chantier"], -(x["updated"] or 0)))

    # ---- par ressource : RAF de son profil face à sa capacité prévue restante
    raf: dict[str, dict[str, float]] = defaultdict(lambda: {"dev": 0.0, "qa": 0.0, "n": 0})
    unassigned = {"dev": 0.0, "qa": 0.0, "n": 0}
    for r in rows:
        sd = history.side_of(r)
        if sd not in SIDES or r.get("done"):
            continue
        tgt = raf[r["assignee"]] if r.get("assignee") else unassigned
        tgt[sd] += r.get(f"raf_{sd}") or 0
        tgt["n"] += 1
    people, seen = [], set()
    for p in (tp or {}).get("people", []):
        name = p["name"]
        mine = raf.get(name, {"dev": 0.0, "qa": 0.0, "n": 0})
        prof = p.get("profile")
        if prof not in SIDES:  # TL / PO sans reste à faire : pas de charge de réalisation
            if mine["dev"] + mine["qa"] < 0.005:
                continue
            prof = "qa" if mine["qa"] > mine["dev"] else "dev"
        aff = p.get("affectation") or {}
        capa = p.get("capa") if p.get("in_team") else None
        if chantier:
            share = (aff.get(chantier) or 0) / sum(aff.values()) if sum(aff.values()) else 0.0
            if share <= 0 and mine["dev"] + mine["qa"] < 0.005:
                continue  # ni affectée au chantier, ni chargée dessus
            capa = None if capa is None else capa * share
        seen.add(name)
        people.append(_person(name, prof, capa, mine, aff, p.get("role", ""), p.get("in_team", False), bool(chantier)))
    for name, mine in raf.items():
        if name not in seen and mine["dev"] + mine["qa"] > 0.005:
            prof = "qa" if mine["qa"] > mine["dev"] else "dev"
            people.append(_person(name, prof, None, mine, {}, "", False))
    order = {"red": 0, "amber": 1, "info": 2, "green": 3, "none": 4}
    people.sort(key=lambda x: (x["profile"], order[x["state"]["lvl"]], -(x["rate"] or 0), x["name"]))

    tot = {sd: {"charge": _r(sum(p["charge"] for p in people if p["profile"] == sd) + unassigned[sd]),
                "capa": _r(sum(p["capa"] or 0 for p in people if p["profile"] == sd)),
                "unassigned": _r(unassigned[sd])} for sd in SIDES}
    open_tasks = [r for r in rows if history.side_of(r) in SIDES and not r.get("done")]
    ready = {
        "open": len(open_tasks),
        "unassigned": unassigned["n"],
        "noest": sum(1 for r in open_tasks if not (r.get("estim_hj") or 0) and not (r.get("raf_hj") or 0)),
    }
    return {
        "sprint": meta.get("sprint"), "today": meta.get("today"), "chantier": chantier, "chantiers": chantiers,
        "milestones": {"dev": meta.get("freeze_code"), "qa": meta.get("fin_test"), "start": meta.get("sprint_start")},
        "test_pct": test_pct, "feats": out_feats, "people": people, "totals": tot, "ready": ready,
        "tempo_error": tempo_error,
        "rules": {"macro_red": MACRO_OVER_RED, "test_tol": TEST_TOL, "test_red": TEST_RED,
                  "load_high": LOAD_HIGH, "load_tense": LOAD_TENSE, "load_low": LOAD_LOW},
    }


def _person(name: str, prof: str, capa: float | None, mine: dict, aff: dict, role: str, in_team: bool, by_ch: bool = False) -> dict[str, Any]:
    charge, other = mine[prof], mine["qa" if prof == "dev" else "dev"]
    rate = charge / capa if capa and capa > 0.05 else None
    if charge < 0.005:
        st = {"lvl": "info", "word": "sans charge"} if capa and capa > 0.05 else {"lvl": "none", "word": "—"}
    elif rate is None:
        st = {"lvl": "red", "word": "non affecté au chantier" if in_team and by_ch else "sans capacité" if in_team else "hors équipe Tempo"}
    elif rate > LOAD_HIGH:
        st = {"lvl": "red", "word": "surcharge"}
    elif rate >= LOAD_TENSE:
        st = {"lvl": "amber", "word": "tendu"}
    elif rate < LOAD_LOW:
        st = {"lvl": "info", "word": "sous-chargé"}
    else:
        st = {"lvl": "green", "word": "équilibré"}
    return {"name": name, "profile": prof, "role": role, "in_team": in_team, "capa": _r(capa), "charge": _r(charge),
            "other": _r(other), "gap": _r(None if capa is None else capa - charge), "rate": _r(None if rate is None else 100 * rate, 1),
            "affectation": aff, "state": st}


def messages(a: dict[str, Any]) -> dict[str, list[str]]:
    """Messages clés de chaque analyse (synthèse de la page et du rapport)."""
    f = a["feats"]
    out: dict[str, list[str]] = {"macro": [], "test": [], "load": [], "ready": []}
    missing = [x for x in f if x["macro"] is None]
    over = sorted((x for x in f if x["macro_state"]["word"] == "dépassement"), key=lambda x: -(x["gap"] or 0))
    if over:
        out["macro"].append(f"{len(over)} FEAT au-delà du macro-chiffrage : " + ", ".join(
            f"{x['key']} (+{jh(x['gap'])} JH, {pct(x['gap_pct'], True)})" for x in over[:3]) + ("…" if len(over) > 3 else ""))
    elif f and len(missing) < len(f):
        out["macro"].append("Toutes les FEAT chiffrées restent dans leur macro-chiffrage.")
    if missing:
        out["macro"].append(f"Macro-chiffrage à saisir pour {len(missing)} FEAT sur {len(f)}.")
    tp = a["test_pct"]
    low = [x for x in f if x["test_state"]["word"] == "test sous-dimensionné"]
    high = [x for x in f if x["test_state"]["word"] == "test sur-dimensionné"]
    dev = sum(x["dev"]["fin"] for x in f)
    qa = sum(x["qa"]["fin"] for x in f)
    if dev > 0.005:
        out["test"].append(f"Test = {round(100 * qa / dev)} % du développement sur le sprint (hypothèse {round(tp)} %).")
    if low:
        out["test"].append(f"Test sous-dimensionné : " + ", ".join(f"{x['key']} ({round(x['test_ratio'])} %)" for x in low[:3]) + ("…" if len(low) > 3 else ""))
    if high:
        out["test"].append(f"Test sur-dimensionné : " + ", ".join(f"{x['key']} ({round(x['test_ratio'])} %)" for x in high[:3]) + ("…" if len(high) > 3 else ""))
    ppl = a["people"]
    red = [p for p in ppl if p["state"]["lvl"] == "red"]
    free = [p for p in ppl if p["state"]["word"] in ("sous-chargé", "sans charge")]
    if red:
        out["load"].append(f"Arbitrage : " + ", ".join(f"{p['name']} ({p['state']['word']}{', ' + pct(p['rate']) if p['rate'] is not None else ''})" for p in red[:4]) + ("…" if len(red) > 4 else ""))
    if free:
        out["load"].append(f"Capacité disponible : " + ", ".join(f"{p['name']} ({jh(p['gap'])} JH)" for p in free[:4]) + ("…" if len(free) > 4 else ""))
    for sd in SIDES:
        t = a["totals"][sd]
        if t["charge"] or t["capa"]:
            out["load"].append(f"{PROF[sd]} : {jh(t['charge'])} JH de charge pour {jh(t['capa'])} JH de capacité ({'+' if t['capa'] - t['charge'] >= 0 else ''}{jh(t['capa'] - t['charge'])} JH).")
    if a["tempo_error"]:
        out["load"].insert(0, f"Capacité indisponible (Tempo) : {a['tempo_error']}")
    rd = a["ready"]
    if rd["unassigned"]:
        out["ready"].append(f"{rd['unassigned']} tâche(s) non affectée(s) : {jh(a['totals']['dev']['unassigned'])} JH DEV, {jh(a['totals']['qa']['unassigned'])} JH QA hors comparatif par ressource.")
    if rd["noest"]:
        out["ready"].append(f"{rd['noest']} tâche(s) ouverte(s) sans estimation : chiffrage actualisé incomplet.")
    if not out["ready"]:
        out["ready"].append(f"Toutes les tâches ouvertes ({rd['open']}) sont estimées et affectées.")
    return out


# ---------------------------------------------------------------- rapport PowerPoint
LVL = {"red": (RED, RED_PALE), "amber": (AMBER_TXT, AMBER_PALE), "green": (OK_TXT, GREEN_PALE), "info": (QA_BLUE, BLUE_PALE),
       "unknown": (MUTED, PANEL), "none": (MUTED, None)}
H = {"bold": True, "color": WHITE, "fill": GREEN_DARK, "align": PP_ALIGN.RIGHT}
HL = {**H, "align": PP_ALIGN.LEFT}
R = {"align": PP_ALIGN.RIGHT}
ROWS = 14  # lignes de tableau par diapositive


def _state_cell(st: dict) -> tuple[str, dict]:
    fg, bg = LVL[st["lvl"]]
    return st["word"], {"bold": True, "color": fg, "fill": bg, "align": PP_ALIGN.CENTER}


def _sub(a: dict) -> str:
    ms = a["milestones"]
    parts = ([f"Chantier {a['chantier']}"] if a.get("chantier") else []) + [f"Situation au {long_date(a['today'])}"]
    if ms.get("dev"):
        parts.append(f"Freeze Code {short_date(ms['dev'])}")
    if ms.get("qa"):
        parts.append(f"fin de test {short_date(ms['qa'])}")
    return " · ".join(parts)


def _cut(s: str, n: int = 64) -> str:
    """Libellé raccourci sans couper un mot."""
    return s if len(s) <= n else s[:n].rsplit(" ", 1)[0].rstrip(" ,;:·-") + "…"


def _short_name(name: str) -> str:
    """« Mohamed Cherif HAMDI » → « Mohamed Cherif H. » (graphiques)."""
    parts = name.split()
    return " ".join(parts[:-1] + [parts[-1][0] + "."]) if len(parts) > 1 and parts[-1].isupper() else name


def _sgn(x: float | None) -> str:
    return ("+" if (x or 0) > 0.005 else "") + jh(x)


def _chunks(xs: list, n: int = ROWS) -> list[list]:
    return [xs[i:i + n] for i in range(0, len(xs), n)] or [[]]


def _summary(prs, n, a, msg):
    sl = _slide(prs, "31_Title Slide")
    _title(sl, f"Cadrage du sprint {a['sprint']}" + (f" – {a['chantier']}" if a.get("chantier") else " – synthèse"), _sub(a))
    cards = [("1. Macro-chiffrage vs chiffrage actualisé", msg["macro"] or ["Aucune FEAT dans le sprint."]),
             ("2. Charge de test vs développement", msg["test"] or ["Aucune charge de développement."]),
             ("3. Charge vs capacité par ressource", msg["load"] or ["Aucune ressource chargée."]),
             ("Préalable : estimation et affectation", msg["ready"])]
    w, h = 6.35, 2.75
    for i, (head, lines) in enumerate(cards):
        x, y = 0.25 + (i % 2) * (w + 0.2), 1.05 + (i // 2) * (h + 0.2)
        _text(sl, x, y, w, h, [[(head, {"bold": True, "size": 15, "color": GREEN_DARK})], [("", {"size": 6})]] + [[("• " + m, {"size": 13})] for m in lines[:5]],
              fill=PANEL, margin=0.18)
        bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
        bar.fill.solid(); bar.fill.fore_color.rgb = GREEN_DARK
        bar.line.fill.background(); bar.shadow.inherit = False
    _page_no(sl, n)
    return n + 1


def _macro(prs, n, a):
    f = a["feats"]
    pages = _chunks(f)
    for i, page in enumerate(pages, 1):
        sl = _slide(prs, "31_Title Slide")
        _title(sl, "1. Macro-chiffrage initial vs chiffrage actualisé" + (f" ({i}/{len(pages)})" if len(pages) > 1 else ""),
               "Chiffrage actualisé = consommé + reste à faire des tâches DEV et QA de la FEAT · en JH")
        rows = [[("Chantier", HL), ("FEAT", HL), ("Macro", H), ("Actualisé", H), ("dont DEV", H), ("dont QA", H), ("Écart", H), ("Écart %", H), ("État", {**H, "align": PP_ALIGN.CENTER})]]
        for x in page:
            over = x["macro_state"]["word"] == "dépassement"
            c = {"color": RED if over else INK, "bold": over, **R}
            rows.append([(x["chantier"], {"size": 9}), (_cut(f"{x['key']} · {x['name']}"), {"size": 9, "bold": over}),
                         (jh(x["macro"]) if x["macro"] is not None else "à saisir", {**R, "color": MUTED if x["macro"] is None else INK}),
                         (jh(x["updated"]), {**R, "bold": True}), (jh(x["dev"]["fin"]), {**R, "color": MUTED}), (jh(x["qa"]["fin"]), {**R, "color": MUTED}),
                         (("+" if (x["gap"] or 0) > 0.005 else "") + jh(x["gap"]), c), (pct(x["gap_pct"], True), c), _state_cell(x["macro_state"])])
        if i == len(pages) and len(f) > 1:
            m = [x["macro"] for x in f if x["macro"] is not None]
            up = sum(x["updated"] for x in f if x["macro"] is not None)
            rows.append([("Total (FEAT chiffrées)", {"bold": True, "fill": PANEL}), ("", {"fill": PANEL}), (jh(sum(m)) if m else "—", {**R, "bold": True, "fill": PANEL}),
                         (jh(up) if m else "—", {**R, "bold": True, "fill": PANEL}), ("", {"fill": PANEL}), ("", {"fill": PANEL}),
                         (("+" if m and up - sum(m) > 0.005 else "") + jh(up - sum(m)) if m else "—", {**R, "bold": True, "fill": PANEL}), ("", {"fill": PANEL}), ("", {"fill": PANEL})])
        widths = [1.7, 3.9, 0.85, 0.95, 0.85, 0.85, 0.85, 0.85, 1.4] if len(page) > 6 else [1.5, 2.8, 0.75, 0.85, 0.75, 0.75, 0.75, 0.75, 1.35]
        _table(sl, 0.2, 1.05, widths, rows, row_h=0.36, size=10)
        if len(page) <= 6:  # place pour le graphique à droite
            chart_feats = [x for x in page if x["macro"] is not None] or page
            _chart(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, 10.6, 1.0, 2.65, 4.2, [x["key"] for x in chart_feats],
                   [("Macro", [x["macro"] or 0 for x in chart_feats]), ("Actualisé", [x["updated"] for x in chart_feats])],
                   "Macro vs actualisé (JH)", [GREY, GREEN_DARK], labels=True, fmt="0.0", gap=60)
        _text(sl, 0.2, 6.75, 12.8, 0.3, [f"Dépassement : actualisé supérieur au macro-chiffrage (rouge au-delà de +{round(a['rules']['macro_red'])} %) · macro-chiffrage saisi dans l'outil (page Cadrage)"],
              size=9, color=MUTED)
        _page_no(sl, n); n += 1
    return n


def _test(prs, n, a):
    f = [x for x in a["feats"] if x["dev"]["fin"] > 0.005 or x["qa"]["fin"] > 0.005]
    tp = a["test_pct"]
    pages = _chunks(f)
    for i, page in enumerate(pages, 1):
        sl = _slide(prs, "31_Title Slide")
        _title(sl, "2. Charge de test vs charge de développement" + (f" ({i}/{len(pages)})" if len(pages) > 1 else ""),
               f"Hypothèse : test = {round(tp)} % du développement · charge = consommé + reste à faire · en JH")
        rows = [[("Chantier", HL), ("FEAT", HL), ("DEV", H), ("QA", H), ("QA ÷ DEV", H), (f"QA théorique ({round(tp)} %)", H), ("Écart QA", H), ("État", {**H, "align": PP_ALIGN.CENTER})]]
        for x in page:
            bad = x["test_state"]["lvl"] in ("red", "amber")
            rows.append([(x["chantier"], {"size": 9}), (_cut(f"{x['key']} · {x['name']}"), {"size": 9}), (jh(x["dev"]["fin"]), R), (jh(x["qa"]["fin"]), {**R, "bold": True}),
                         (pct(x["test_ratio"]) if x["test_ratio"] is not None else "—", {**R, "bold": True, "color": LVL[x["test_state"]["lvl"]][0] if bad else INK}),
                         (jh(x["test_theo"]), {**R, "color": MUTED}), (("+" if (x["test_gap"] or 0) > 0.005 else "") + jh(x["test_gap"]), R), _state_cell(x["test_state"])])
        if i == len(pages) and len(f) > 1:
            dev, qa = sum(x["dev"]["fin"] for x in f), sum(x["qa"]["fin"] for x in f)
            P = {"fill": PANEL}
            rows.append([("Total", {"bold": True, **P}), ("", P), (jh(dev), {**R, "bold": True, **P}), (jh(qa), {**R, "bold": True, **P}),
                         (pct(100 * qa / dev) if dev > 0.005 else "—", {**R, "bold": True, **P}), (jh(tp / 100 * dev), {**R, **P}),
                         (("+" if qa - tp / 100 * dev > 0.005 else "") + jh(qa - tp / 100 * dev), {**R, "bold": True, **P}), ("", P)])
        narrow = len(page) <= 6
        widths = [1.5, 2.9, 0.8, 0.8, 0.9, 1.15, 0.85, 1.6] if narrow else [1.7, 4.4, 0.9, 0.9, 1.0, 1.35, 1.0, 1.6]
        _table(sl, 0.2, 1.05, widths, rows, row_h=0.36, size=10)
        if narrow:
            _chart(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, 10.75, 1.0, 2.5, 4.2, [x["key"] for x in page],
                   [("QA réel", [x["qa"]["fin"] for x in page]), (f"QA théorique {round(tp)} %", [x["test_theo"] for x in page])],
                   "Charge QA (JH)", [QA_BLUE, GREY], labels=True, fmt="0.0", gap=60)
        _text(sl, 0.2, 6.75, 12.8, 0.3, [f"Conforme à ±{round(a['rules']['test_tol'])} points de l'hypothèse ; fort écart au-delà de {round(a['rules']['test_red'])} points · "
                                         "test sous-dimensionné : risque de recette incomplète ; sur-dimensionné : estimation de test à revoir"], size=9, color=MUTED)
        _page_no(sl, n); n += 1
    return n


def _load(prs, n, a):
    ppl = [p for p in a["people"] if p["state"]["lvl"] != "none"]
    if a["tempo_error"] and not ppl:
        sl = _slide(prs, "31_Title Slide")
        _title(sl, "3. Charge planifiée vs capacité disponible", _sub(a))
        dnf._no_data(sl, 0.5, 2.0, 12.3, 2.0, "Capacité indisponible", a["tempo_error"])
        _page_no(sl, n)
        return n + 1
    blocks = []
    for sd in SIDES:
        mine = [p for p in ppl if p["profile"] == sd]
        if mine:
            blocks.append(("head", sd))
            blocks += [("p", p) for p in mine if p["state"]["word"] != "sans charge"]
            free = [p for p in mine if p["state"]["word"] == "sans charge"]
            if free:
                blocks.append(("free", free))
    pages = _chunks(blocks, 15)
    for k in range(len(pages) - 1):  # pas d'en-tête de profil seul en bas de page
        if pages[k] and pages[k][-1][0] == "head":
            pages[k + 1].insert(0, pages[k].pop())
    for i, page in enumerate(pages, 1):
        sl = _slide(prs, "31_Title Slide")
        _title(sl, "3. Charge planifiée vs capacité disponible par ressource" + (f" ({i}/{len(pages)})" if len(pages) > 1 else ""),
               "Charge = reste à faire des tâches affectées (DEV pour un DEV, QA pour un QA) · capacité prévue restante jusqu'au jalon (Tempo) · en JH")
        rows = [[("Ressource", HL), ("Capacité", H), ("Charge", H), ("Écart", H), ("Charge ÷ capa", H), ("Hors profil", H), ("État", {**H, "align": PP_ALIGN.CENTER})]]
        for kind, x in page:
            if kind == "head":
                fill = GREEN_PALE if x == "dev" else BLUE_PALE
                t = a["totals"][x]
                rows.append([(f"{PROF[x]} · jusqu'au {short_date(a['milestones'][x])}", {"bold": True, "fill": fill}), (jh(t["capa"]), {**R, "bold": True, "fill": fill}),
                             (jh(t["charge"]), {**R, "bold": True, "fill": fill}), (_sgn((t["capa"] or 0) - (t["charge"] or 0)), {**R, "bold": True, "fill": fill}),
                             (pct(100 * t["charge"] / t["capa"]) if t["capa"] else "—", {**R, "bold": True, "fill": fill}),
                             (f"non affecté {jh(t['unassigned'])}" if t["unassigned"] else "", {**R, "fill": fill, "color": RED if t["unassigned"] else MUTED, "size": 9}), ("", {"fill": fill})])
                continue
            if kind == "free":
                fg, bg = LVL["info"]
                rows.append([(f"Sans charge ({len(x)}) : " + ", ".join(p["name"] for p in x), {"size": 9, "color": fg}),
                             (jh(sum(p["capa"] or 0 for p in x)), R), ("0,00", {**R, "bold": True}), (_sgn(sum(p["capa"] or 0 for p in x)), {**R, "bold": True, "color": OK_TXT}),
                             ("0 %", {**R, "bold": True, "color": fg}), ("", {}), ("sans charge", {"bold": True, "color": fg, "fill": bg, "align": PP_ALIGN.CENTER})])
                continue
            fg = LVL[x["state"]["lvl"]][0]
            rows.append([(x["name"] + ("" if x["in_team"] else " (hors Tempo)"), {}), (jh(x["capa"]), R), (jh(x["charge"]), {**R, "bold": True}),
                         (_sgn(x["gap"]), {**R, "bold": True, "color": RED if (x["gap"] or 0) < -0.005 else OK_TXT}),
                         (pct(x["rate"]) if x["rate"] is not None else "—", {**R, "bold": True, "color": fg}),
                         (jh(x["other"], blank_zero=True), {**R, "color": MUTED}), _state_cell(x["state"])])
        # ligne « sans charge » : la liste des noms peut prendre deux lignes
        _table(sl, 0.2, 1.05, [3.6, 0.95, 0.95, 0.95, 1.15, 1.0, 1.4], rows, row_h=0.33, size=10)
        # graphique : capacité vs charge des ressources de la page
        pp = [x for k, x in page if k == "p"]
        if pp:
            _chart(sl, XL_CHART_TYPE.BAR_CLUSTERED, 10.35, 1.0, 2.9, min(5.6, 1.2 + 0.4 * len(pp)), [_short_name(x["name"]) for x in pp][::-1],
                   [("Capacité", [x["capa"] or 0 for x in pp][::-1]), ("Charge", [x["charge"] for x in pp][::-1])],
                   "Capacité vs charge (JH)", [GREY, GREEN_DARK], labels=False, fmt="0.0", gap=40)
        _text(sl, 0.2, 6.75, 12.8, 0.3, ["Surcharge > 100 % : arbitrage · tendu ≥ 90 % · sous-chargé < 60 % ou sans charge : rééquilibrage possible · "
                                         "hors profil : RAF de l'autre côté (non compté) · un TL / PO sans reste à faire n'apparaît pas"], size=9, color=MUTED)
        _page_no(sl, n); n += 1
    return n


def build_pptx(a: dict[str, Any]) -> bytes:
    prs = Presentation(str(dnf.TEMPLATE))
    dnf._fonts(prs)
    cov = _slide(prs, "Diapositive de titre heka", keep=(0, 1))
    cov.placeholders[0].text = "Cadrage du sprint"
    sub = cov.placeholders[1].text_frame
    sub.text = f"Estimation, affectation et répartition du reste à faire · sprint {a['sprint']}" + (f" · {a['chantier']}" if a.get("chantier") else "")
    sub.add_paragraph().text = f"Le {short_date(a['today'])}/{(a['today'] or '')[:4]}"
    if dnf.LOGO.exists():
        cov.shapes.add_picture(str(dnf.LOGO), Emu(10172418), Emu(265078), Emu(1900601), Emu(792476))
    msg = messages(a)
    n = _summary(prs, 2, a, msg)
    n = _macro(prs, n, a)
    n = _test(prs, n, a)
    n = _load(prs, n, a)
    end = _slide(prs, "31_Title Slide", keep=(0,))
    end.placeholders[0].text = "Merci pour votre attention"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
