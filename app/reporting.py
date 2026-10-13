"""Calculs de la vue « Reporting Tasks » — reproduit les colonnes calculées de la feuille Charge."""
from __future__ import annotations

from datetime import date
from typing import Any

from .config import Settings
from .models import Dataset, Task


def chantier_of(labels: list[str], s: Settings) -> str:
    """Même logique que la formule SI(ESTNUM(CHERCHE(...))) : premier motif trouvé, insensible à la casse."""
    text = ", ".join(labels).lower()
    for c in s.chantiers:
        if str(c["match"]).lower() in text:
            return c["name"]
    return ""


def type_charge(t: str, s: Settings) -> str:
    tl = t.lower()
    if tl in s.types_ano:
        return "ANO"
    if tl in s.types_bug:
        return "BUG"
    if tl in s.types_qa:
        return "QA"
    if tl in s.types_dev:
        return "DEV"
    return ""


def _hj(hours: float | None, s: Settings) -> float:
    return round((hours or 0) / s.hours_per_day, 4)


RAF_FINI_LABEL = "RAF à solder dans Jira"


def solde_finies(rows: list[dict]) -> list[dict]:
    """Une tâche finie n'a plus de reste à faire : son RAF Jira passe dans raf_residuel (alerte « à solder »)
    et ne compte dans aucun total. Appliqué aussi aux relevés enregistrés avant cette règle (13/10/2026)."""
    for r in rows:
        if r.get("done") and "raf_residuel" not in r:
            r["raf_residuel"] = r.get("raf_hj") or 0.0
            r["raf_hj"] = r["raf_dev"] = r["raf_qa"] = 0.0
        for a in r.get("alerts") or ():
            if a.get("code") == "RAF_FINI":
                a["label"] = RAF_FINI_LABEL
    return rows


def _alerts(t: Task, done: bool, is_dev: bool, is_qa: bool, raf_hj: float, s: Settings, today: date) -> list[dict]:
    out: list[dict] = []
    add = lambda code, level, label: out.append({"code": code, "level": level, "label": label})  # noqa: E731
    if done:
        if raf_hj > 0:
            add("RAF_FINI", "info", RAF_FINI_LABEL)
        return out
    # Livrée (Ready for QA, Merge Request) : ni retard ni RAF à estimer, le travail de la tâche est fait
    delivered = t.status.lower() in s.no_delay_statuses
    if t.status.lower() == "blocked":
        add("BLOQUEE", "danger", "Bloquée")
    if not t.due:
        add("SANS_ECHEANCE", "warn", "Sans échéance")
    else:
        if t.due < today and not delivered:
            add("RETARD", "danger", f"En retard ({(today - t.due).days} j)")
        if is_dev and s.freeze_code and t.due > s.freeze_code:
            add("APRES_FREEZE", "danger", "Échéance après Freeze Code")
        if is_qa and s.fin_test and t.due > s.fin_test:
            add("APRES_FIN_TEST", "danger", "Échéance après fin de test")
    # Aucune estimation (ni initiale ni restante) sur une tâche à faire ou en cours : charge invisible dans le RAF
    no_estimate = t.original_h in (None, 0) and t.remaining_h in (None, 0)
    if no_estimate and t.status.lower() in s.estimate_required_statuses:
        add("SANS_ESTIMATION", "warn", "Sans estimation")
    elif (t.remaining_h in (None, 0)) and (t.spent_h or 0) > 0 and not delivered:
        add("RAF_A_ESTIMER", "warn", "Travaux non finalisés, RAF à estimer")
    if not t.assignee:
        add("SANS_ASSIGNE", "warn", "Non assignée")
    return out


# Sens d'un lien Jira (libellé anglais ou français de l'instance) → libellé affiché
_REL = {"blocks": "Bloque", "bloque": "Bloque", "is blocked by": "Bloqué par", "est bloqué par": "Bloqué par",
        "relates to": "Lié à", "relates_to": "Lié à", "se rapporte à": "Lié à", "is caused by": "Causé par", "causes": "Cause"}


def _rel(kind: str) -> str:
    return _REL.get(kind.strip().lower(), kind.strip() or "Lié à")


def build_rows(ds: Dataset, s: Settings, browse_base: str, today: date | None = None) -> list[dict[str, Any]]:
    today = today or date.today()
    rows = []
    # Toutes les US du sprint liées à chaque tâche (liens vus depuis l'US ou depuis la tâche) :
    # un bug qui bloque plusieurs US les cite toutes
    us_of: dict[str, set[str]] = {}
    for st in ds.stories.values():
        for k in st.linked_keys:
            us_of.setdefault(k, set()).add(st.key)
    for t in ds.tasks:
        us = ds.stories.get(t.us_key)
        linked_us = sorted(us_of.get(t.key, set()) | {k for k in t.linked_keys if k in ds.stories} | ({t.us_key} if t.us_key else set()))
        # FEAT : celle de l'US, sinon celle portée par la tâche (bug de régression rattaché directement à la FEAT)
        feat_key = (us.epic_key if us and us.epic_key else "") or t.epic_key
        feat = ds.features.get(feat_key) if feat_key else None
        tl = t.type.lower()
        is_dev, is_qa = tl in s.raf_dev_types, tl in s.raf_qa_types
        done = t.status.lower() in s.done_statuses
        raf_hj = _hj(t.remaining_h, s)
        spent_hj, estim_hj = _hj(t.spent_h, s), _hj(t.original_h, s)

        if done:
            avancement: float | None = 1.0
        elif (t.spent_h or 0) > 0 and (t.remaining_h or 0) > 0:
            avancement = t.spent_h / (t.spent_h + t.remaining_h)
        elif (t.spent_h or 0) > 0:
            avancement = None  # RAF à estimer
        else:
            avancement = 0.0

        rows.append({
            "key": t.key,
            "url": f"{browse_base}/browse/{t.key}" if browse_base else "",
            "type": t.type,
            "type_charge": type_charge(t.type, s),
            "summary": t.summary,
            "status": t.status,
            "done": done,
            "delivered": not done and t.status.lower() in s.no_delay_statuses,
            "assignee": t.assignee,
            "due": t.due.isoformat() if t.due else None,
            "due_dev": t.due.isoformat() if (is_dev and t.due) else None,
            "due_qa": t.due.isoformat() if (is_qa and t.due) else None,
            "raf_hj": raf_hj,
            "raf_dev": raf_hj if is_dev else 0.0,
            "raf_qa": raf_hj if is_qa else 0.0,
            # côté de la tâche pour l'effort (planifié vs consommé) : dev, qa, ou None (bogues…)
            "side": "dev" if is_dev else "qa" if is_qa else None,
            "spent_hj": spent_hj,
            "estim_hj": estim_hj,
            "avancement": avancement,
            "chantier": chantier_of(t.labels, s) or "(Sans chantier)",
            "sprint": t.sprint,
            "project": t.project,
            "labels": t.labels,
            "updated": t.updated.isoformat() if t.updated else None,
            # suivi des anomalies : arrivée, résolution (délai de correction) et priorité
            "created": t.created.isoformat() if t.created else None,
            "resolved": t.resolved.isoformat() if t.resolved else None,
            "priority": t.priority,
            "us_key": t.us_key,
            "us_url": f"{browse_base}/browse/{t.us_key}" if (browse_base and t.us_key and ds.source == "jira") else "",
            "us_summary": us.summary if us else ("" if not t.us_key else t.us_key),
            "us_status": us.status if us else "",
            # US du sprint liées, avec le sens du lien (« Bloque », « Lié à »…) pour la liste des bugs
            "linked_us": [{"key": k, "summary": ds.stories[k].summary, "rel": _rel(t.link_kinds.get(k, "")),
                           "url": f"{browse_base}/browse/{k}" if (browse_base and ds.source == "jira") else ""}
                          for k in linked_us if k in ds.stories],
            "feat": (feat.summary if feat else feat_key) or "(Sans FEAT)",
            "feat_key": feat_key,
            "feat_src": ("us" if us and us.epic_key else "tache") if feat_key else "",
            "feat_status": feat.status if feat else "",
            "alerts": _alerts(t, done, is_dev, is_qa, raf_hj, s, today),
        })
    _deduce_feat(rows, ds, s)
    return solde_finies(rows)


def _deduce_feat(rows: list[dict[str, Any]], ds: Dataset, s: Settings) -> None:
    """Tâches et bugs sans FEAT dans Jira : FEAT par défaut du chantier (config), sinon la FEAT unique du chantier."""
    by_ch: dict[str, set[str]] = {}
    for r in rows:
        if r["feat_key"]:
            by_ch.setdefault(r["chantier"], set()).add(r["feat_key"])
    for r in rows:
        if r["feat_key"]:
            continue
        key = s.feat_par_chantier.get(r["chantier"]) or (next(iter(by_ch[r["chantier"]])) if len(by_ch.get(r["chantier"], ())) == 1 else "")
        if not key:
            continue
        feat = ds.features.get(key)
        r.update(feat=feat.summary if feat else key, feat_key=key, feat_src="chantier", feat_status=feat.status if feat else "")
        r["alerts"].append({"code": "FEAT_DEDUITE", "level": "info", "label": "FEAT déduite du chantier : à rattacher dans Jira"})


def filter_rows(rows: list[dict], params: dict[str, str]) -> list[dict]:
    """Filtres identiques à ceux de l'écran (utilisés pour l'export)."""
    def ok(r: dict) -> bool:
        for field in ("chantier", "feat", "assignee", "status", "type_charge"):
            v = params.get(field)
            if v and r.get(field) != v:
                return False
        q = (params.get("q") or "").lower()
        if q and q not in f"{r['key']} {r['summary']} {r['us_summary']} {r['assignee']}".lower():
            return False
        if params.get("alerts") == "1" and not [a for a in r["alerts"] if a["level"] != "info"]:
            return False
        if params.get("hideDone") == "1" and r["done"]:
            return False
        return True
    return [r for r in rows if ok(r)]


def to_xlsx(rows: list[dict], meta: dict) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporting Tasks"
    ws.append([f"Reporting Tasks — {meta.get('sprint', '')}", "", f"Freeze Code : {meta.get('freeze_code') or '-'}",
               f"Fin de test : {meta.get('fin_test') or '-'}", f"Extraction : {meta.get('generated_at', '')}"])
    ws.append([])
    cols = [("Chantier", "chantier"), ("FEAT mère", "feat"), ("US mère", "us_key"), ("Résumé US", "us_summary"),
            ("Statut US mère", "us_status"), ("Clé", "key"), ("Type", "type"), ("Résumé", "summary"),
            ("État", "status"), ("Personne assignée", "assignee"), ("Date échéance DEV", "due_dev"),
            ("Date échéance QA", "due_qa"), ("RAF DEV (JH)", "raf_dev"), ("RAF QA (JH)", "raf_qa"),
            ("RAF à solder, tâche finie (JH)", "raf_residuel"), ("Alertes", "alerts")]
    ws.append([c[0] for c in cols])
    head = ws.max_row
    for c in ws[head]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="15803D")  # vert de la charte be-ivs
        c.alignment = Alignment(vertical="center", wrap_text=True)
    from datetime import date as _d
    for r in sorted(rows, key=lambda x: (x["chantier"], x["us_summary"], x["type"], x["key"])):
        line = []
        for _, k in cols:
            v = r.get(k)
            if k == "alerts":
                v = ", ".join(a["label"] for a in v)
            elif k in ("due_dev", "due_qa") and v:
                v = _d.fromisoformat(v)
            line.append(v)
        ws.append(line)
        for k in ("due_dev", "due_qa"):
            ws.cell(ws.max_row, [c[1] for c in cols].index(k) + 1).number_format = "DD/MM/YYYY"
    widths = [22, 30, 13, 40, 16, 13, 18, 45, 15, 20, 14, 14, 11, 11, 14, 35]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(head + 1, 1)
    ws.auto_filter.ref = f"A{head}:{get_column_letter(len(cols))}{ws.max_row}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
